import io
import zipfile
import pytest
import pptx

from sovereign.core.state.models import Task, TaskStatus, Finding, Decision, EvidenceReference, UnresolvedQuestion
from sovereign.core.artifacts.models import ArtifactRequest, ArtifactType
from sovereign.infrastructure.artifacts.renderers import PptxRenderer


def test_pptx_renderer_valid_presentation():
    task = Task(title='Turbine T-501 Inspection', goal='Assess vibration anomaly and bearing wear', status=TaskStatus.COMPLETED)
    ev1 = EvidenceReference(source_id='doc_vibration_log', locator='Slide 4', metadata={'text': 'Peak amplitude 4.2 mm/s.'})
    f1 = Finding(statement='Bearing degradation detected in radial plane.', confidence='high', evidence_refs=[ev1.evidence_id])
    d1 = Decision(decision='SCHEDULE_MAINTENANCE', rationale='Vibration exceeds ISO Category C.')
    q1 = UnresolvedQuestion(question='Verify spare bearing availability.', status='OPEN')

    req = ArtifactRequest(
        task_id=task.task_id,
        artifact_type=ArtifactType.PPTX,
        title='Executive Briefing: Turbine T-501',
        requested_sections=['executive_summary', 'findings', 'evidence', 'decisions', 'unresolved_questions', 'limitations']
    )

    renderer = PptxRenderer()
    pptx_bytes = renderer.render(task, [f1, d1, q1], [ev1], req)

    assert isinstance(pptx_bytes, bytes)
    assert len(pptx_bytes) > 0

    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as zf:
        assert '[Content_Types].xml' in zf.namelist()
        assert 'ppt/presentation.xml' in zf.namelist()

    prs = pptx.Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 5

    title_slide = prs.slides[0]
    assert 'Executive Briefing: Turbine T-501' in title_slide.shapes.title.text

    findings_slide = prs.slides[2]
    assert '2. Technical Findings' in findings_slide.shapes.title.text
    full_findings_text = ' '.join(p.text for s in findings_slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs)
    assert 'Bearing degradation detected' in full_findings_text
    assert 'doc_vibration_log:Slide 4' in full_findings_text


def test_pptx_renderer_empty_sections():
    task = Task(title='Empty Presentation', goal='No state items')
    req = ArtifactRequest(
        task_id=task.task_id,
        artifact_type=ArtifactType.PPTX,
        title='Empty Briefing',
        requested_sections=['executive_summary', 'findings', 'evidence', 'decisions', 'limitations']
    )

    renderer = PptxRenderer()
    pptx_bytes = renderer.render(task, [], [], req)

    prs = pptx.Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 5
    findings_slide = prs.slides[2]
    full_findings_text = ' '.join(p.text for s in findings_slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs)
    assert 'No technical findings recorded' in full_findings_text
