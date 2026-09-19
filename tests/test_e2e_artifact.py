"""End-to-End tests for Artifact Generation."""
import pytest
from pathlib import Path
from sovereign.core.state.models import Task, Finding, EvidenceReference
from sovereign.core.artifacts.models import ArtifactRequest, ArtifactType, ArtifactStatus
from sovereign.infrastructure.artifacts.storage import LocalArtifactStorage
from sovereign.infrastructure.artifacts.engine import LocalArtifactEngine
from sovereign.infrastructure.state.sqlite_repository import SQLiteTaskRepository
from sovereign.infrastructure.knowledge.sqlite_knowledge import SQLiteKnowledgeBase
from sovereign.infrastructure.knowledge.chunkers.structure_aware_chunker import StructureAwareChunker
from sovereign.infrastructure.config import get_settings

def test_e2e_inspection_artifact(tmp_path, monkeypatch):
    """
    Simulates the end-to-end flow:
    1. Knowledge Base indexing (Pkg 03)
    2. Task State holding evidence/findings (Pkg 02)
    3. Artifact Engine reading state and generating markdown (Pkg 06)
    """
    settings = get_settings()
    monkeypatch.setattr(settings, "artifact_dir_name", str(tmp_path / "artifacts"))
    
    # 1. Pkg 03 - Knowledge Base
    kb = SQLiteKnowledgeBase(db_path=tmp_path / "kb.db")
    
    from sovereign.core.knowledge.models import Document, DocumentStatus
    doc = Document(filename="inspection_report.txt", source_path="inspection_report.txt", document_type="text", file_size=100, content_hash="mock_hash", status=DocumentStatus.INGESTED)
    kb.register_document(doc)
    doc_id = doc.document_id
    from sovereign.infrastructure.state.tokenizer import ApproximateTokenCounter
    tokenizer = ApproximateTokenCounter()
    chunker = StructureAwareChunker(tokenizer=tokenizer)
    from sovereign.core.knowledge.models import DocumentBlock
    block = DocumentBlock(document_id=doc_id, page_number=37, sequence=0, text="CRITICAL FAILURE: Valve X ruptured")
    chunks = chunker.chunk(doc, [block])
    kb.save_chunks(chunks)
    
    # 2. Pkg 02 - Task State (Simulating AgentHost output)
    repo = SQLiteTaskRepository(db_path=tmp_path / "state.db")
    task = Task(title="Valve Inspection", goal="Inspect valve status")
    repo.create_task(task)
    
    ev = EvidenceReference(source_id=doc_id, locator="37", metadata={"chunk_id": chunks[0].chunk_id, "text": chunks[0].text})
    repo.add_evidence(task.task_id, ev)
    
    finding = Finding(statement="Valve X has ruptured according to the report.", confidence="high", evidence_refs=[ev.evidence_id])
    repo.add_state_item(task.task_id, finding)
    
    # 3. Pkg 06 - Artifact Engine
    storage = LocalArtifactStorage()
    engine = LocalArtifactEngine(repo, storage)
    
    req = ArtifactRequest(
        task_id=task.task_id,
        artifact_type=ArtifactType.MARKDOWN,
        title="Final Inspection Summary",
        requested_sections=["executive_summary", "findings", "evidence"]
    )
    
    artifact = engine.generate(req)
    
    # Validations
    assert artifact.status == ArtifactStatus.COMPLETED
    assert len(artifact.source_references) == 1
    assert artifact.source_references[0].locator == "37"
    
    filepath = Path(artifact.metadata.file_path)
    assert filepath.exists()
    content = filepath.read_text(encoding="utf-8")
    
    assert "Final Inspection Summary" in content
    assert "Valve X has ruptured according to the report." in content
    assert doc_id in content
    assert "37" in content
    
    assert len(content) < 1000  # Artifact should be bounded and concise


def test_e2e_docx_inspection_artifact(tmp_path, monkeypatch):
    """
    Simulates the end-to-end flow generating a DOCX deliverable:
    1. Knowledge Base indexing (Pkg 03)
    2. Task State holding evidence/findings (Pkg 02)
    3. Artifact Engine reading state and generating .docx (Pkg 06)
    4. Cryptographic SHA-256 hash validation
    """
    import hashlib
    import docx
    settings = get_settings()
    monkeypatch.setattr(settings, "artifact_dir_name", str(tmp_path / "artifacts"))
    
    # 1. Pkg 03 - Knowledge Base
    kb = SQLiteKnowledgeBase(db_path=tmp_path / "kb_docx.db")
    
    from sovereign.core.knowledge.models import Document, DocumentStatus
    doc = Document(filename="inspection_v204.txt", source_path="inspection_v204.txt", document_type="text", file_size=120, content_hash="mock_hash_docx", status=DocumentStatus.INGESTED)
    kb.register_document(doc)
    doc_id = doc.document_id
    from sovereign.infrastructure.state.tokenizer import ApproximateTokenCounter
    tokenizer = ApproximateTokenCounter()
    chunker = StructureAwareChunker(tokenizer=tokenizer)
    from sovereign.core.knowledge.models import DocumentBlock
    block = DocumentBlock(document_id=doc_id, page_number=2, sequence=0, text="ANOMALY: Flange V-204 seal wear 0.15mm exceeds tolerance")
    chunks = chunker.chunk(doc, [block])
    kb.save_chunks(chunks)
    
    # 2. Pkg 02 - Task State
    repo = SQLiteTaskRepository(db_path=tmp_path / "state_docx.db")
    task = Task(title="V-204 Flange Analysis", goal="Analyze flange anomaly report and extract seal measurements")
    repo.create_task(task)
    
    ev = EvidenceReference(source_id=doc_id, locator="Page 2", metadata={"chunk_id": chunks[0].chunk_id, "text": chunks[0].text})
    repo.add_evidence(task.task_id, ev)
    
    finding = Finding(statement="Flange V-204 seal wear measured at 0.15mm (critical limit 0.10mm).", confidence="high", evidence_refs=[ev.evidence_id])
    repo.add_state_item(task.task_id, finding)
    
    # 3. Pkg 06 - Artifact Engine
    storage = LocalArtifactStorage()
    engine = LocalArtifactEngine(repo, storage)
    
    req = ArtifactRequest(
        task_id=task.task_id,
        artifact_type=ArtifactType.DOCX,
        title="Industrial Analysis Note: Flange V-204",
        requested_sections=["executive_summary", "findings", "evidence", "limitations"]
    )
    
    artifact = engine.generate(req)
    
    # Validations
    assert artifact.status == ArtifactStatus.COMPLETED
    assert artifact.type == ArtifactType.DOCX
    assert len(artifact.source_references) == 1
    assert artifact.source_references[0].locator == "Page 2"
    
    filepath = Path(artifact.metadata.file_path)
    assert filepath.exists()
    assert filepath.suffix == ".docx"
    
    # Validate binary content and SHA-256
    file_bytes = filepath.read_bytes()
    expected_hash = hashlib.sha256(file_bytes).hexdigest()
    assert artifact.metadata.content_hash == expected_hash
    assert artifact.metadata.size_bytes == len(file_bytes)
    
    # Load and verify DOCX content
    doc_parsed = docx.Document(filepath)
    full_text = "\n".join(p.text for p in doc_parsed.paragraphs)
    assert "Industrial Analysis Note: Flange V-204" in full_text
    assert "Flange V-204 seal wear measured at 0.15mm" in full_text
    assert "doc_inspection_v204" in full_text or doc_id in full_text
    assert "Page 2" in full_text


def test_e2e_xlsx_and_pptx_artifacts(tmp_path, monkeypatch):
    """Verify E2E generation and SHA-256 byte comparison for XLSX and PPTX artifacts."""
    import hashlib
    import zipfile
    import pptx
    settings = get_settings()
    monkeypatch.setattr(settings, "artifact_dir_name", str(tmp_path / "artifacts"))

    repo = SQLiteTaskRepository(db_path=tmp_path / "state_multi.db")
    task = Task(title="Compressor C-301 Audit", goal="Audit compressor discharge pressure and generate multi-format deliverables")
    repo.create_task(task)

    ev = EvidenceReference(source_id="doc_press_01", locator="Table 2", metadata={"text": "Pressure 18.2 bar exceeds 16.0 bar."})
    repo.add_evidence(task.task_id, ev)

    finding = Finding(statement="Discharge pressure 18.2 bar exceeds operating rating.", confidence="high", evidence_refs=[ev.evidence_id])
    repo.add_state_item(task.task_id, finding)

    storage = LocalArtifactStorage()
    engine = LocalArtifactEngine(repo, storage)

    # 1. Generate XLSX
    req_xlsx = ArtifactRequest(
        task_id=task.task_id,
        artifact_type=ArtifactType.XLSX,
        title="Pressure Log: Compressor C-301",
        requested_sections=["executive_summary", "findings", "evidence"]
    )
    art_xlsx = engine.generate(req_xlsx)
    assert art_xlsx.status == ArtifactStatus.COMPLETED
    assert art_xlsx.type == ArtifactType.XLSX
    xlsx_path = Path(art_xlsx.metadata.file_path)
    assert xlsx_path.exists()
    xlsx_bytes = xlsx_path.read_bytes()
    assert hashlib.sha256(xlsx_bytes).hexdigest() == art_xlsx.metadata.content_hash

    # Verify XLSX structure
    with zipfile.ZipFile(xlsx_path) as zf:
        assert "xl/workbook.xml" in zf.namelist()

    # 2. Generate PPTX
    req_pptx = ArtifactRequest(
        task_id=task.task_id,
        artifact_type=ArtifactType.PPTX,
        title="Executive Briefing: Compressor C-301",
        requested_sections=["executive_summary", "findings", "evidence", "limitations"]
    )
    art_pptx = engine.generate(req_pptx)
    assert art_pptx.status == ArtifactStatus.COMPLETED
    assert art_pptx.type == ArtifactType.PPTX
    pptx_path = Path(art_pptx.metadata.file_path)
    assert pptx_path.exists()
    pptx_bytes = pptx_path.read_bytes()
    assert hashlib.sha256(pptx_bytes).hexdigest() == art_pptx.metadata.content_hash

    # Verify PPTX structure
    prs = pptx.Presentation(pptx_path)
    assert len(prs.slides) == 5

    # 3. Verify SHA-256 byte comparison detects tampering
    tampered_bytes = xlsx_bytes + b"\x00TAMPERED"
    tampered_hash = hashlib.sha256(tampered_bytes).hexdigest()
    assert tampered_hash != art_xlsx.metadata.content_hash

